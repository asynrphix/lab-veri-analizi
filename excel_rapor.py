import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import LineChart, Reference, Series

df = pd.read_csv("analiz_sonucu.csv")

wb = Workbook()
ws = wb.active
ws.title = "Veri"

for satir in dataframe_to_rows(df, index=False, header=True):
    ws.append(satir)

for hucre in ws[1]:
    hucre.font = Font(bold=True)

kirmizi = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
limit_kolon_no = df.columns.get_loc("Limit_Asimi") + 1

for satir_no in range(2, ws.max_row + 1):
    deger = ws.cell(row=satir_no, column=limit_kolon_no).value
    if deger == True:
        for kolon_no in range(1, ws.max_column + 1):
            ws.cell(row=satir_no, column=kolon_no).fill = kirmizi

for kolon in ws.columns:
    max_uzunluk = max(len(str(hucre.value)) for hucre in kolon)
    ws.column_dimensions[kolon[0].column_letter].width = max_uzunluk + 2


# --- ÖZET SAYFASI ---
ozet = wb.create_sheet("Özet")

ortalama = df["% Silica Concentrate"].mean()
std_sapma = df["% Silica Concentrate"].std()
asan_sayisi = df["Limit_Asimi"].sum()
toplam_gun = len(df)

ozet["A1"] = "Laboratuvar Analiz Özeti"
ozet["A1"].font = Font(bold=True, size=14)

ozet["A3"] = "Toplam gün sayısı:"
ozet["B3"] = toplam_gun
ozet["A4"] = "Ortalama % Silica Concentrate:"
ozet["B4"] = round(ortalama, 2)
ozet["A5"] = "Standart sapma:"
ozet["B5"] = round(std_sapma, 2)
ozet["A6"] = "Hesaplanan limit (ort + 1.5×sapma):"
ozet["B6"] = round(ortalama + 1.5 * std_sapma, 2)
ozet["A7"] = "Limiti aşan gün sayısı:"
ozet["B7"] = int(asan_sayisi)
ozet["A8"] = "Limiti aşan gün oranı:"
ozet["B8"] = f"%{round(asan_sayisi / toplam_gun * 100, 1)}"

for satir in range(3, 9):
    ozet[f"A{satir}"].font = Font(bold=True)

ozet.column_dimensions["A"].width = 35
ozet.column_dimensions["B"].width = 15

# --- GRAFİK ---
grafik = LineChart()
grafik.title = "% Silica Concentrate - Günlük Trend"
grafik.x_axis.title = "Gün"
grafik.y_axis.title = "% Silica Concentrate"
grafik.width = 24
grafik.height = 10

silica_kolon_no = df.columns.get_loc("% Silica Concentrate") + 1
degerler = Reference(ws, min_col=silica_kolon_no, min_row=2, max_row=ws.max_row)
seri = Series(degerler, title="% Silica Concentrate")
grafik.series.append(seri)

kategoriler = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
grafik.set_categories(kategoriler)

ozet.add_chart(grafik, "D25")  # AI yorumunun altına, çakışmasın diye aşağı aldık

# --- AI YORUMU ---
try:
    with open("ai_yorumu.txt", "r", encoding="utf-8") as f:
        ai_metin = f.read()
except FileNotFoundError:
    ai_metin = "AI yorumu henüz üretilmedi. Önce gemini_yorum.py çalıştırılmalı."

ozet["A11"] = "AI Değerlendirmesi:"
ozet["A11"].font = Font(bold=True, size=12)
ozet["A12"] = ai_metin
ozet["A12"].alignment = ozet["A12"].alignment.copy(wrap_text=True)
ozet.merge_cells("A12:H20")
ozet.row_dimensions[12].height = 150

wb.save("lab_raporu_final.xlsx")
print("Kaydedildi: lab_raporu_final.xlsx")